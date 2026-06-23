from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable


try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIRS = {
    "probepruefung": ROOT / "probepruefungen",
    "lernzettel": ROOT / "lernzettel",
}
OUTPUT_DIR = ROOT / "_material_index"

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".docx",
    ".xlsx",
    ".csv",
}


@dataclass
class MaterialEntry:
    kind: str
    path: str
    extension: str
    size_bytes: int
    sha256: str
    pages_or_sheets: int | None
    text_chars: int | None
    extraction_status: str
    note: str


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def extract_pdf_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = len(reader.pages)
        chars = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            chars += len(text)
        status = "ok" if chars else "no_text_or_scanned"
        note = "PDF text layer found." if chars else "No PDF text layer detected; OCR may be needed."
        return pages, chars, status, note
    except Exception as exc:  # pragma: no cover - defensive for unknown PDFs
        return None, None, "error", f"{type(exc).__name__}: {exc}"


def extract_docx_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    try:
        from docx import Document

        doc = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                table_text.append("\t".join(cell.text for cell in row.cells))
        chars = len(text) + len("\n".join(table_text))
        return None, chars, "ok", "DOCX text extracted for statistics."
    except Exception as exc:  # pragma: no cover
        return None, None, "error", f"{type(exc).__name__}: {exc}"


def extract_xlsx_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        chars = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        chars += len(str(value))
        return len(workbook.worksheets), chars, "ok", "XLSX values read for statistics."
    except Exception as exc:  # pragma: no cover
        return None, None, "error", f"{type(exc).__name__}: {exc}"


def extract_text_file_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = path.read_text(encoding=encoding)
            return None, len(text), "ok", f"Text read as {encoding}."
        except UnicodeDecodeError:
            continue
    return None, None, "error", "Could not decode text file as UTF-8 or CP1252."


def extract_csv_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            chars = 0
            rows = 0
            with path.open("r", encoding=encoding, newline="") as handle:
                for row in csv.reader(handle):
                    rows += 1
                    chars += sum(len(cell) for cell in row)
            return rows, chars, "ok", f"CSV read as {encoding}; pages_or_sheets means rows."
        except UnicodeDecodeError:
            continue
    return None, None, "error", "Could not decode CSV as UTF-8 or CP1252."


def file_stats(path: Path) -> tuple[int | None, int | None, str, str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_stats(path)
    if suffix == ".docx":
        return extract_docx_stats(path)
    if suffix == ".xlsx":
        return extract_xlsx_stats(path)
    if suffix in {".txt", ".md"}:
        return extract_text_file_stats(path)
    if suffix == ".csv":
        return extract_csv_stats(path)
    return None, None, "unsupported", "Unsupported extension."


def scan_files() -> Iterable[tuple[str, Path]]:
    for kind, directory in INPUT_DIRS.items():
        directory.mkdir(exist_ok=True)
        for path in sorted(directory.rglob("*")):
            if path.is_file():
                yield kind, path


def build_inventory() -> list[MaterialEntry]:
    entries: list[MaterialEntry] = []
    for kind, path in scan_files():
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            pages, chars, status, note = None, None, "unsupported", "Unsupported extension."
        else:
            pages, chars, status, note = file_stats(path)
        entries.append(
            MaterialEntry(
                kind=kind,
                path=rel(path),
                extension=extension or "(none)",
                size_bytes=path.stat().st_size,
                sha256=sha256_file(path),
                pages_or_sheets=pages,
                text_chars=chars,
                extraction_status=status,
                note=note,
            )
        )
    return entries


def write_jsonl(entries: list[MaterialEntry], path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for entry in entries:
            handle.write(json.dumps(asdict(entry), ensure_ascii=False) + "\n")


def write_markdown(entries: list[MaterialEntry], path: Path) -> None:
    lines = [
        "# AP1-Materialinventar",
        "",
        "Dieses Inventar wird automatisch aus `probepruefungen/` und `lernzettel/` erzeugt.",
        "Es speichert keine Rohtexte geschützter Prüfungsaufgaben.",
        "",
        "| Art | Datei | Typ | Größe | Seiten/Blätter | Textzeichen | Status | Hinweis |",
        "|---|---|---|---:|---:|---:|---|---|",
    ]
    for entry in entries:
        pages = "" if entry.pages_or_sheets is None else str(entry.pages_or_sheets)
        chars = "" if entry.text_chars is None else str(entry.text_chars)
        note = entry.note.replace("|", "/")
        lines.append(
            f"| {entry.kind} | `{entry.path}` | {entry.extension} | {entry.size_bytes} | "
            f"{pages} | {chars} | {entry.extraction_status} | {note} |"
        )
    if not entries:
        lines.append("| - | - | - | - | - | - | leer | Keine Dateien gefunden. |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--jsonl", default="_material_index/material_inventory.jsonl")
    parser.add_argument("--md", default="_material_index/AP1_MATERIAL_INVENTAR.md")
    args = parser.parse_args()

    OUTPUT_DIR.mkdir(exist_ok=True)
    entries = build_inventory()
    jsonl_path = ROOT / args.jsonl
    md_path = ROOT / args.md
    jsonl_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.parent.mkdir(parents=True, exist_ok=True)
    write_jsonl(entries, jsonl_path)
    write_markdown(entries, md_path)

    print(f"Dateien gefunden: {len(entries)}")
    print(f"JSONL: {rel(jsonl_path)}")
    print(f"Markdown: {rel(md_path)}")
    for entry in entries:
        print(f"- {entry.kind}: {entry.path} ({entry.extraction_status})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
